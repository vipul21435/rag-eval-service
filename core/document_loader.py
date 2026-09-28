"""Document loader: text extraction for the supported file formats.

PDF, TXT and Markdown work with the base install. DOCX, PPTX and Excel need
the ``documents`` extra. Problems are reported as exceptions with a message
meant for the user: ``UnsupportedFormatError`` for an extension nothing here
can read, ``MissingParserError`` for a format whose optional parser is not
installed. The ingestion pipeline catches them per file, so a bad file never
aborts a batch.
"""

from __future__ import annotations

import logging
import os
from io import StringIO

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = frozenset({".txt", ".md"})
EXCEL_EXTENSIONS = frozenset({".xlsx", ".xls"})
SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".pptx"} | TEXT_EXTENSIONS | EXCEL_EXTENSIONS)


class UnsupportedFormatError(ValueError):
    """The file extension is not one of ``SUPPORTED_EXTENSIONS``."""


class MissingParserError(RuntimeError):
    """The format is supported but its optional parser is not installed."""


def describe_supported_formats() -> str:
    """Comma-separated list of the supported extensions, for error messages."""
    return ", ".join(sorted(SUPPORTED_EXTENSIONS))


def extract_text(filepath: str | os.PathLike[str]) -> str:
    """Return the plain-text content of ``filepath``.

    Raises ``UnsupportedFormatError`` for an unknown extension,
    ``MissingParserError`` when the format needs the ``documents`` extra,
    and ``ValueError`` for a text file that is not UTF-8.
    """
    file_ext = os.path.splitext(os.fspath(filepath))[1].lower()

    if file_ext == ".pdf":
        return _extract_pdf(filepath)
    if file_ext in TEXT_EXTENSIONS:
        return _extract_utf8(filepath)
    if file_ext == ".docx":
        return _extract_docx(filepath)
    if file_ext in EXCEL_EXTENSIONS:
        return _extract_excel(filepath)
    if file_ext == ".pptx":
        return _extract_pptx(filepath)

    raise UnsupportedFormatError(
        f"unsupported file format {file_ext or '(none)'!r}; supported: {describe_supported_formats()}"
    )


def _extract_utf8(filepath: str | os.PathLike[str]) -> str:
    try:
        with open(filepath, encoding="utf-8") as file:
            return file.read()
    except UnicodeDecodeError as exc:
        raise ValueError("text file is not UTF-8 encoded") from exc


def _extract_pdf(filepath: str | os.PathLike[str]) -> str:
    from pdfminer.high_level import extract_text_to_fp

    output = StringIO()
    with open(filepath, "rb") as file:
        extract_text_to_fp(file, output)
    return output.getvalue()


def _missing_parser(package: str, what: str) -> MissingParserError:
    return MissingParserError(f"reading {what} requires {package}; install the 'documents' extra")


def _extract_docx(filepath: str | os.PathLike[str]) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise _missing_parser("python-docx", ".docx files") from exc
    document = Document(os.fspath(filepath))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_excel(filepath: str | os.PathLike[str]) -> str:
    try:
        import pandas as pd
    except ImportError as exc:
        raise _missing_parser("pandas", "Excel files") from exc
    parts: list[str] = []
    workbook = pd.ExcelFile(filepath)
    for sheet_name in workbook.sheet_names:
        frame = workbook.parse(sheet_name)
        parts.append(f"Sheet: {sheet_name}\n{frame.to_string(index=False)}\n\n")
    return "".join(parts)


def _extract_pptx(filepath: str | os.PathLike[str]) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise _missing_parser("python-pptx", ".pptx files") from exc
    presentation = Presentation(os.fspath(filepath))
    lines: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text:
                lines.append(text)
    return "\n".join(lines) + ("\n" if lines else "")
